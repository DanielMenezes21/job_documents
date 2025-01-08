from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from telas.contrato_hono.texto_clausula_adv import ADVOGADOS_OAB, TEXTOS_CLAUSULA1, TEXTOS_CLAUSULA3
from docx import Document
import os
from telas.contrato_hono.extracao_contrato import formatar_data , substituir_palavras_documento
from telas.homepage.home_screen import HomeScreen

def voltar(self, instance):
    """
    Função para voltar para a tela anterior ou a tela principal.
    """
    self.manager.current = "processo_screen"
    
def atualizar_dados(screen_instance, dados):
    """
    Atualiza os dados recebidos na tela Processo2Screen.
    """
    screen_instance.dados = dados
    screen_instance.caminho_modelo = dados.get("caminho_modelo", None)  # Armazena o caminho do arquivo modelo
    print(f"Dados recebidos: {screen_instance.dados}")

def on_text_selected(self, modelo_spinner, text):
    """
    Carrega o texto selecionado no TextInput.
    """
    if text in TEXTOS_CLAUSULA1:
        self.text_input.text = TEXTOS_CLAUSULA1[text]  # Preenche o TextInput com o texto selecionado
        
def on_text_selected2(self, modelo_spinner2, text):
    """
    Carrega o texto selecionado no TextInput.
    """
    if text in TEXTOS_CLAUSULA3:
        self.text_input2.text = TEXTOS_CLAUSULA3[text]  # Preenche o TextInput com o texto selecionado

def salvar_texto(screen_instance, _):
    """
    Salva o texto editado e atualiza o documento.
    """
    try:
        
        caminho_modelo = os.path.join(os.path.dirname(__file__), "14 CRIMINAL -  CONT, PROC - teste.docx")
        
        if not screen_instance.dados:
            mostrar_popup(screen_instance, "Erro", "Nenhum dado foi recebido da tela inicial!")
            return

        # Obtém o texto editado
        clausula1 = screen_instance.text_input.text.strip()
        clausula3 = screen_instance.text_input2.text.strip()
        nome_arquivo = screen_instance.nome_arquivo_input.text.strip() + '.docx'

        if not screen_instance.caminho_modelo:
            mostrar_popup(screen_instance, "Erro", "O arquivo modelo não foi selecionado na tela inicial.")
            return

        # Verifica se o arquivo existe
        if not os.path.exists(screen_instance.caminho_modelo):
            mostrar_popup(screen_instance, "Erro", f"O arquivo {screen_instance.caminho_modelo} não foi encontrado!")
            return

        # Obter o nome do advogado OAB baseado no nome selecionado no spinner
        advogado_nome = screen_instance.adv_spinner.text
        advogado_oab = ADVOGADOS_OAB.get(advogado_nome, "OAB não encontrado")
        data_atual = formatar_data()
        # Dicionário de placeholders
        placeholders = {
            "Hd1ad": clausula1,
            "9has91": clausula3,
            "hdas98": advogado_oab,
            "#NOME_CONTRATANTE": screen_instance.dados.get("nome_contratante"),
            "#CONTRATANTE_CPF": screen_instance.dados.get("cpf"),
            "#RG": screen_instance.dados.get("rg"),
            "#SEC_RG": screen_instance.dados.get("sec_rg"),
            "dh91a": screen_instance.dados.get("cidade_contratante"),
            "=SECESTCONT": screen_instance.dados.get("sigla_estado_contratante"),
            "%INSTRA": screen_instance.dados.get("inscrita_o"),
            "#NACIONALIDADE": screen_instance.dados.get("nacionalidade"),
            "#SIGLA_ESTADO": screen_instance.dados.get("est_rg_input"),
            "#DATA_AGORA": data_atual,
        }
        exit_user_file = screen_instance.caminho_modelo
        substituir_palavras_documento(screen_instance.caminho_modelo,placeholders, nome_arquivo)
    
        mostrar_popup(screen_instance, "Sucesso", "Documento salvo com sucesso!")
        
    except Exception as e:
        print(f"Erro ao obter dados: {e}")
        return None
    

def mostrar_popup(screen_instance, titulo, mensagem):
    """
    Exibe um popup com uma mensagem.
    """
    conteudo = BoxLayout(orientation='vertical', padding=10, spacing=10)
    conteudo.add_widget(Label(text=mensagem, halign='center'))
    btn_fechar = Button(text='Fechar', size_hint=(1, 0.3))
    popup = Popup(title=titulo, content=conteudo, size_hint=(0.7, 0.4))
    btn_fechar.bind(on_press=popup.dismiss)
    conteudo.add_widget(btn_fechar)
    popup.open()